#!/usr/bin/env python3
"""Generate OOP lab 2 report DOCX from lab 1 template."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parent
OOP = ROOT.parent
TEMPLATE = OOP / "ооп_лаба1_отчет.docx"
OUT = OOP / "ооп_лаба2_отчет.docx"
UML_PNG = ROOT / "report_assets" / "uml_class_diagram.png"
OBJ_EMPTY_PNG = ROOT / "report_assets" / "object_state1_empty.png"
OBJ_THREE_PNG = ROOT / "report_assets" / "object_state2_three_records.png"
ACTIVITY_INSERT_PNG = ROOT / "report_assets" / "activity_insert.png"
VTABLE_DEMO_CPP = OOP / "vtable_demo" / "vtable_demo.cpp"


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def delete_table(table) -> None:
    element = table._element
    element.getparent().remove(element)


def set_run_font(run, size_pt=14, bold=False, italic=False, name="Times New Roman") -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_line_spacing_15(paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def style_body_paragraph(paragraph, first_line=True) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    set_line_spacing_15(paragraph)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Cm(1.25) if first_line else Cm(0)


def add_body(doc: Document, text: str, first_line=True) -> None:
    p = doc.add_paragraph()
    style_body_paragraph(p, first_line=first_line)
    run = p.add_run(text)
    set_run_font(run)


def add_list_item(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Paragraph")
    style_body_paragraph(p, first_line=False)
    run = p.add_run("— " + text)
    set_run_font(run)


def add_heading1(doc: Document, text: str, center=False) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    set_line_spacing_15(p)
    run = p.add_run(text)
    set_run_font(run, size_pt=16, bold=True)


def add_heading2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_line_spacing_15(p)
    run = p.add_run(text)
    set_run_font(run, size_pt=14, bold=True)


def add_caption_table(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.first_line_indent = Cm(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run)


def add_caption_figure(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.first_line_indent = Cm(0)
    set_line_spacing_15(p)
    run = p.add_run(text)
    set_run_font(run)


def add_code_block(doc: Document, text: str) -> None:
    for line in text.splitlines():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt = p.paragraph_format
        fmt.first_line_indent = Cm(0)
        fmt.left_indent = Cm(0.5)
        set_line_spacing_15(p)
        run = p.add_run(line if line else " ")
        set_run_font(run, size_pt=11, name="Courier New")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            set_line_spacing_15(p)
            for r in p.runs:
                set_run_font(r, bold=True)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val
            for p in table.rows[ri].cells[ci].paragraphs:
                set_line_spacing_15(p)
                for r in p.runs:
                    set_run_font(r)


def apply_line_spacing_everywhere(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        set_line_spacing_15(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    set_line_spacing_15(paragraph)


def main() -> None:
    doc = Document(str(TEMPLATE))

    # Title page tweaks
    doc.paragraphs[14].text = "ЛАБОРАТОРНАЯ РАБОТА №2"
    doc.paragraphs[16].text = "«Рефакторинг программы»"

    while len(doc.tables) > 0:
        delete_table(doc.tables[0])

    while len(doc.paragraphs) > 43:
        delete_paragraph(doc.paragraphs[43])

    def page_break() -> None:
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    add_heading1(doc, "СОДЕРЖАНИЕ", center=True)
    toc_lines = [
        "ЦЕЛЬ И ЗАДАЧИ РАБОТЫ",
        "ХОД РАБОТЫ",
        "1 Исходный проект",
        "2 Недостатки исходной реализации",
        "3 Проектирование объектной архитектуры",
        "4 Программная реализация",
        "5 Unit-тестирование",
        "ДОПОЛНИТЕЛЬНЫЕ ЗАДАНИЯ",
        "1 UML-диаграммы объектов",
        "2 UML-диаграмма активностей",
        "3 Подмена виртуального метода",
        "ВЫВОДЫ",
    ]
    for line in toc_lines:
        p = doc.add_paragraph(line)
        style_body_paragraph(p, first_line=False)
        if line.startswith(("1 ", "2 ", "3 ", "4 ", "5 ")):
            p.paragraph_format.left_indent = Cm(1.0)
        for run in p.runs:
            set_run_font(run)

    page_break()
    add_heading1(doc, "ЦЕЛЬ И ЗАДАЧИ РАБОТЫ", center=True)

    add_body(
        doc,
        "Цель работы: освоение этапов анализа предметной области и проектирования "
        "архитектуры программы в объектной модели на примере рефакторинга процедурного проекта.",
    )
    add_body(doc, "Задачи работы:")
    add_list_item(doc, "проанализировать исходную программу и выделить ключевые абстракции предметной области;")
    add_list_item(doc, "выявить недостатки архитектуры и качества кода исходного проекта;")
    add_list_item(doc, "спроектировать объектную архитектуру с UML-диаграммой классов;")
    add_list_item(doc, "реализовать новую архитектуру на языке C++;")
    add_list_item(doc, "разработать unit-тесты для основного функционала.")
    add_body(
        doc,
        "Выполненные дополнительные задания: №1 (UML-диаграммы объектов), "
        "№2 (UML-диаграмма активностей), №3 (подмена виртуального метода и vtable).",
    )

    add_body(doc, "Этапы выполнения:")
    add_list_item(doc, "анализ и проектирование (UML-диаграмма классов);")
    add_list_item(doc, "рефакторинг исходного кода и разработка unit-тестов.")

    page_break()
    add_heading1(doc, "ХОД РАБОТЫ", center=True)

    add_heading2(doc, "1 Исходный проект")
    add_body(
        doc,
        "В качестве исходного проекта использована программа lab_tree (язык C) — "
        "консольная телефонная книга с индексом на сбалансированном AVL-дереве. "
        "Предметная область: хранение записей «телефонный номер — имя абонента» и выполнение "
        "операций insert, find, delete, tree по текстовому протоколу из файла.",
    )
    add_body(
        doc,
        "Программа читает команды из input.txt и пишет ответы в output.txt. "
        "Номер задаётся в кавычках в формате 8XXXXXXXXXX (11 символов) или +7XXXXXXXXXX "
        "(12 символов); имя — строка в кавычках с поддержкой экранирования. "
        "Ключом дерева служат десять цифр номера; при обходе tree выводится каноническое "
        "представление номера с префиксом 8 или +7.",
    )
    add_body(
        doc,
        "Исходный код (~460 строк) объединяет разбор команд, логику базы, AVL-операции и "
        "ручное управление памятью (malloc/free) в наборе глобальных функций и структур "
        "Record, Tree, Database.",
    )

    add_heading2(doc, "2 Недостатки исходной реализации")
    add_body(doc, "К недостаткам архитектуры и качества кода относятся:")
    add_list_item(
        doc,
        "процедурная организация: отсутствуют классы, ответственность смешана "
        "(функция parsing одновременно разбирает строки и управляет базой);",
    )
    add_list_item(
        doc,
        "слабая инкапсуляция: поля структур доступны напрямую, инварианты дерева "
        "не защищены границами объектов;",
    )
    add_list_item(
        doc,
        "ручное управление памятью для Record и узлов дерева, риск ошибок при изменениях;",
    )
    add_list_item(
        doc,
        "коды ошибок (KEY_EXISTS, NOT_FOUND) и строки протокола разбросаны по функциям;",
    )
    add_list_item(
        doc,
        "низкая тестируемость: нельзя изолированно проверить разбор номера или вставку в дерево "
        "без полного файлового сценария.",
    )

    add_heading2(doc, "3 Проектирование объектной архитектуры")
    add_body(
        doc,
        "При декомпозиции выделены абстракции: запись абонента (Record и вложенный Number), "
        "AVL-дерево (AVLTree), база данных (Database), разбор команд (Parser), "
        "исключение предметной области (AVLTreeError). Класс Number инкапсулирует проверку "
        "формата номера; AVLTree отвечает только за структуру данных; Database — за сценарии "
        "insert/find/del/tree; Parser — за синтаксис команд и текст ответов.",
    )
    add_body(
        doc,
        "Связи между классами (нотация UML.pdf): Parser ассоциирован с Database и вызывает его "
        "методы; Database композиционно включает AVLTree; узлы дерева хранят Record; "
        "Record содержит Number; при ошибках Database и конструктор Number бросают AVLTreeError.",
    )

    add_body(doc, "На рисунке 1 представлена UML-диаграмма классов разработанной архитектуры.", first_line=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing_15(p)
    run = p.add_run()
    run.add_picture(str(UML_PNG), width=Cm(14))
    add_caption_figure(doc, "Рисунок 1 — UML-диаграмма классов телефонной книги")

    add_caption_table(doc, "Таблица 1 — Классы объектной архитектуры")
    add_table(
        doc,
        ["Класс", "Назначение"],
        [
            [
                "Record::Number",
                "Валидация строки номера, хранение префикса (8 или +7) и 10 цифр ключа; "
                "при некорректном формате — исключение AVLTreeError.",
            ],
            [
                "Record",
                "Запись телефонной книги: номер и имя абонента (имя хранится в том виде, "
                "в каком передано протоколом, включая кавычки).",
            ],
            [
                "AVLTree",
                "AVL-дерево по ключу number.digits(): вставка, поиск имени, удаление, "
                "форматированный обход tree; балансировка и узлы скрыты внутри класса.",
            ],
            [
                "Database",
                "Фасад предметной области: insert, find, del, tree; владеет AVLTree и "
                "счётчиком size_; преобразует ошибки дерева в исключения key exists / not found.",
            ],
            [
                "Parser",
                "Чтение команд insert/find/delete/tree построчно, разбор кавычек и имён, "
                "формирование строк протокола (insert:success, find:..., incorrect и т.д.).",
            ],
            [
                "AVLTreeError",
                "Тип исключения с сообщением what(); наследует std::exception.",
            ],
        ],
    )

    add_heading2(doc, "4 Программная реализация")
    add_body(
        doc,
        "Рефакторинг выполнен на C++17. Точка входа main.cpp открывает input.txt и output.txt, "
        "создаёт объекты Database и Parser и вызывает parse. Память узлов дерева и записей "
        "управляется через std::unique_ptr; копирование записей при удалении узла локализовано "
        "в методах AVLTree.",
    )
    add_body(
        doc,
        "Поведение протокола сохранено относительно исходной программы и эталонных тестовых "
        "файлов tests/input1.txt … tests/input6.txt и tests/input.txt. Сборка: скрипт build.sh "
        "(исполняемый файл lab2 в каталоге проекта).",
    )

    add_heading2(doc, "5 Unit-тестирование")
    add_body(
        doc,
        "Тесты реализованы на Google Test. Скрипт build_test.sh собирает lab2_test и запускает "
        "все наборы. Unit-тесты проверяют классы Number, Record, AVLTree, Database, Parser и "
        "AVLTreeError; интеграционный набор Lab2Integration сравнивает вывод Parser+Database "
        "с эталонными output*.txt.",
    )
    add_body(
        doc,
        "По результатам прогона ./build_test.sh: 58 тестов из 10 наборов, все успешны (PASSED).",
        first_line=True,
    )

    add_caption_table(doc, "Таблица 2 — Unit-тесты (основные наборы)")
    add_table(
        doc,
        ["Файл / набор", "Проверяемое поведение"],
        [
            ["number_test.cpp", "Конструктор Number: допустимые 8/+7, ошибки формата, текст incorrect number."],
            ["record_test.cpp", "Хранение полей Record для номеров 8 и +7."],
            ["exception_test.cpp", "Сообщения AVLTreeError (key exists, not found, произвольный текст)."],
            ["avl_tree_test.cpp", "Пустое дерево, insert/find, дубликат, remove, formatTree, балансировка."],
            ["database_test.cpp", "Insert/find/del/tree, исключения, вывод tree:(null) и двух узлов."],
            ["parser_test.cpp", "Все команды протокола, incorrect, escape в имени, +7."],
            ["integration_test.cpp", "Сквозные сценарии input1–6 и полный input.txt vs эталоны."],
        ],
    )

    add_body(
        doc,
        "Пример локального прогона: cd oop/lab2 && ./build_test.sh — в конце вывод "
        "[  PASSED  ] 58 tests.",
    )

    page_break()
    add_heading1(doc, "ДОПОЛНИТЕЛЬНЫЕ ЗАДАНИЯ", center=True)

    add_heading2(doc, "1 UML-диаграммы объектов")
    add_body(
        doc,
        "Дополнительное задание №1: построены две UML-диаграммы объектов для двух "
        "состояний программы lab2 после разбора команд из файла (контур Parser + Database). "
        "На диаграммах показаны экземпляры Database, AVLTree, узлов дерева и записей Record "
        "и связи между ними в соответствии с реализацией на std::unique_ptr.",
    )

    add_body(doc, "Состояние 1 — пустая телефонная книга.", first_line=True)
    add_body(
        doc,
        "Момент времени: программа создана, команда tree выполнена на пустой базе "
        "(как первая строка сценария tests/input3.txt). Объект Database существует, "
        "счётчик size_ равен 0, AVLTree не содержит узлов (root_ пуст, empty() истинно). "
        "Объектов Node и Record в памяти нет.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing_15(p)
    p.add_run().add_picture(str(OBJ_EMPTY_PNG), width=Cm(12))
    add_caption_figure(doc, "Рисунок 2 — Диаграмма объектов: пустая база")

    add_body(doc, "Состояние 2 — три записи в AVL-дереве.", first_line=True)
    add_body(
        doc,
        "Момент времени: после трёх успешных команд insert из tests/input3.txt "
        "(номера \"87451002823\", \"86244259344\", \"+75529719054\" и соответствующие имена). "
        "Database.size_ = 3. Топология дерева совпадает с выводом tree после третьей вставки "
        "(корень, левый и правый потомок). Узел Node ссылается на Record; в Record поле number "
        "имеет тип Number (prefix_ и number_, ключ для сравнения в дереве — 10 цифр number_). "
        "Строка name хранит имя в формате протокола. Для трёх insert: numRight (8, 7451002823), "
        "numMid (8, 6244259344), numLeft (7, 5529719054) — соответствуют номерам "
        "\"87451002823\", \"86244259344\", \"+75529719054\".",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing_15(p)
    p.add_run().add_picture(str(OBJ_THREE_PNG), width=Cm(14))
    add_caption_figure(doc, "Рисунок 3 — Диаграмма объектов: три записи")

    add_heading2(doc, "2 UML-диаграмма активностей")
    add_body(
        doc,
        "Дополнительное задание №2: выбран сценарий обработки одной команды insert "
        "функцией handleInsert (файл parser.cpp). Вход — строка вида "
        "insert \"81234567890\" \"Ivan\"; возможные выходы протокола: insert:success, "
        "insert:e_exist, incorrect. Диаграмма построена в одной дорожке (без swim lanes); "
        "для компактной ширины основной поток разбит на два горизонтальных ряда, "
        "ошибочные ветви сходятся к действию incorrect.",
    )
    add_body(
        doc,
        "Порядок проверок соответствует коду: readQuoted для телефона, readName для имени, "
        "контроль лишних символов в аргументах, конструктор Number (ошибка формата даёт incorrect), "
        "Database::insert (дубликат ключа — insert:e_exist через AVLTreeError).",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing_15(p)
    p.add_run().add_picture(str(ACTIVITY_INSERT_PNG), width=Cm(17))
    add_caption_figure(doc, "Рисунок 4 — Диаграмма активностей: команда insert")

    add_heading2(doc, "3 Подмена виртуального метода")
    add_body(
        doc,
        "Дополнительное задание №3: изучение полиморфизма и virtual function table (vtable) "
        "на отдельной тестовой программе (каталог oop/vtable_demo). Модель и версия компилятора: "
        "Apple Clang 17.0.0 (clang-1700.6.3.2), целевая архитектура arm64, macOS. "
        "Отладчик: LLDB (lldb-1703.0.236.21). Сборка: clang++ -g -O0 -std=c++17 vtable_demo.cpp.",
    )
    add_body(
        doc,
        "В программе четыре класса по заданию: интерфейс IWriter (pure virtual write); "
        "абстрактный WriterBase, реализующий write() из интерфейса и объявляющий virtual label() "
        "с реализацией по умолчанию; конкретные WriterA и WriterB с переопределением write() и label(). "
        "Вызов через указатель WriterBase* p на объект WriterA диспетчеризуется через vtable: "
        "в начале объекта хранится указатель vptr на таблицу адресов виртуальных функций; "
        "при p->write() / p->label() код загружает vptr, индексирует vtable и выполняет "
        "функцию по адресу из таблицы (для WriterA — методы WriterA).",
    )
    add_body(
        doc,
        "Подмена для демонстрации выполняется в LLDB: на breakpoint перед виртуальным вызовом "
        "в ячейке vtable, соответствующей методу, записывается адрес реализации WriterBase "
        "(команда memory write -s 8). Автоматизация — скрипт demo.sh (останов, патч, continue). "
        "Сценарий 1: override virtual label() — до патча WriterA::label, после — WriterBase::label. "
        "Сценарий 2: override write() (в IWriter pure virtual, тело в WriterBase) — до патча "
        "WriterA::write, после — WriterBase::write.",
    )
    add_caption_table(doc, "Таблица 3 — Вывод программы до и после подмены vtable")
    add_table(
        doc,
        ["Режим", "Строки вывода (фрагмент для объекта WriterA)"],
        [
            [
                "Обычный запуск (./vtable_demo)",
                "WriterA::write; WriterA::label",
            ],
            [
                "Патч слота label() (demo.sh, блок patch label)",
                "WriterA::write; WriterBase::label",
            ],
            [
                "Патч слота write() (demo.sh, блок patch write)",
                "WriterBase::write; WriterA::label",
            ],
        ],
    )
    add_body(doc, "Исходный код тестовой программы vtable_demo.cpp:", first_line=True)
    add_code_block(doc, VTABLE_DEMO_CPP.read_text(encoding="utf-8"))

    page_break()
    add_heading1(doc, "ВЫВОДЫ", center=True)
    add_body(
        doc,
        "В ходе работы процедурная программа телефонной книги на AVL-дереве преобразована "
        "в объектную архитектуру из шести классов с разделением ответственности между разбором "
        "команд, базой данных и деревом.",
    )
    add_body(
        doc,
        "UML-диаграмма классов отражает композицию Database–AVLTree–Record и роль Parser "
        "как клиента базы. Реализация на C++ с RAII и исключениями упростила сопровождение "
        "по сравнению с ручным malloc/free.",
    )
    add_body(
        doc,
        "Разработанные unit- и интеграционные тесты подтверждают сохранение функционала "
        "исходного проекта; все 58 тестов проходят успешно.",
    )
    add_body(
        doc,
        "Дополнительно выполнены задания №1–№3: UML-диаграммы объектов, UML-диаграмма "
        "активностей для insert и демонстрация vtable с подменой виртуальных методов "
        "в LLDB на программе vtable_demo.",
    )

    apply_line_spacing_everywhere(doc)
    doc.save(str(OUT))
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
